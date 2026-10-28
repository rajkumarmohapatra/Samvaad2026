import csv,io,uuid,zipfile
from openpyxl import load_workbook
from django.core.exceptions import ValidationError,PermissionDenied
from .forms import NominationForm
from .models import Participant,Receipt
from .services import nominate

HEADERS=['name','preferred_name','phone','alternate_phone','email','gender','country','state_or_ut','state','district','village','community','tribal_status','tribe_name','sub_tribe','disability_status','languages','primary_stream','category','source_identifier','attendee_group','element_representative','registration_spoc','registration_mode','remarks']

def preview(user,edition,upload):
    content=upload.read()
    if len(content)>2_000_000:raise ValidationError('Maximum import size is 2 MB.')
    if upload.name.lower().endswith('.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(entry.file_size for entry in archive.infolist())>20_000_000:raise ValidationError('Expanded workbook is too large; split the file.')
        except zipfile.BadZipFile:raise ValidationError('Invalid .xlsx workbook.')
        wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
        try:
            ws=wb['Participants'] if 'Participants' in wb.sheetnames else wb.active
            values=ws.iter_rows(values_only=True);headers=[str(x or '').strip() for x in next(values)]
            if len(headers)>100:raise ValidationError('Too many worksheet columns.')
            raw=[]
            for row in values:
                raw.append(dict(zip(headers,[str(x) if x is not None else '' for x in row])))
                if len(raw)>5000:raise ValidationError('Maximum 5,000 rows per import.')
        finally:wb.close()
    else:
        try:raw=list(csv.DictReader(io.StringIO(content.decode('utf-8-sig'))))
        except (UnicodeDecodeError,csv.Error):raise ValidationError('Use UTF-8 CSV with the supplied headers.')
    if len(raw)>5000:raise ValidationError('Maximum 5,000 rows per import.')
    if raw and set(raw[0])-set(HEADERS):raise ValidationError('Unsupported columns. Use the portal template; identity numbers, photos and calculated helper columns must not be placed in this import.')
    if raw and not set(['name','phone','state','primary_stream','category']).issubset(raw[0]):raise ValidationError('Use the portal template: name,phone,state,primary_stream,category are required.')
    report=[];seen=set()
    for n,row in enumerate(raw,2):
        if not any(str(x).strip() for x in row.values()):continue
        data={str(k).strip():str(v or '').strip() for k,v in row.items() if k is not None}
        status=data.get('tribal_status','')
        data['tribal_status']={'tribal':'Tribal','not tribal':'Not Tribal','non-tribal':'Not Tribal','non tribal':'Not Tribal','not_tribal':'Not Tribal','prefer not to say':'Prefer not to say','not specified':''}.get(status.casefold(),status)
        for field in ['disability_status']:
            value=data.get(field,'');data[field]={'yes':'Yes','no':'No','prefer not to say':'Prefer not to say'}.get(value.casefold(),value)
        data['request_key']=str(uuid.uuid4());data['food_eligible']='on';data['kit_eligible']='on'
        f=NominationForm(data,user=user,edition=edition)
        errors=[]
        if not f.is_valid():errors=[f'{k}: {"; ".join(v)}' for k,v in f.errors.items()]
        mark=(data.get('name','').casefold(),data.get('phone',''))
        if mark in seen:errors.append('Duplicate within file.')
        seen.add(mark)
        if Participant.objects.filter(name__iexact=data.get('name',''),phone=data.get('phone','')).exists():errors.append('Likely existing participant; administrator must resolve identity.')
        report.append({'row':n,'data':data,'errors':errors,'saved':False})
    return report


def commit(user,edition,batch,report):
    if not batch or len(batch)>100:raise ValidationError('Invalid batch key.')
    for row in report:
        key=f'{batch}:{row["row"]}'
        if Receipt.objects.filter(actor=user,key=key).exists():
            # The staged file is server-session-owned. A retry reports prior success.
            row['saved']=True;row['errors']=[];continue
        if row['errors']:continue
        f=NominationForm(row['data'],user=user,edition=edition)
        if not f.is_valid():row['errors']=['Input or permissions changed. Preview again.'];continue
        try:
            nominate(user,edition,key,f.cleaned_data);row['saved']=True
        except (ValidationError,PermissionDenied) as e:row['errors']=[str(e)]
    return report
