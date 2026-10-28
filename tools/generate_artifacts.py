"""Generate inspectable configuration, documentation and synthetic templates."""
import csv, json, sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
UNITS='''Accommodation
Akhra
All Venue Team (Gopal Maidan, TCC etc)
Anchoring - (Cultural) Gopal Maidan
Colors of Jharkhhand
Communications
Crew
Cultural
Cultural Celebrations
Delegate
Delegates Team
Digital & Analytics
Documentation Team
Finance & Accounts
Food Team
Gifts & Kits
Gopal Maidan - Food Stalls
Listening Post
Mobility
Nagara Team
NOTE TAKING
Others
Public Health
Registration & Helpdesk
Rhythms of the Earth
Safety
Samuday Ke Saath
Samvaad Fellowship
Security - Tata Steel
Stakeholders Community
Stakeholders Management
Supply Chain
Team Samvaad
Ticketing & Booking Team
Transport
Tribal Artisans
Tribal Cuisine
Tribal Healers
Tribal Sports
TSL Volunteers'''.splitlines()
def write(path,text):
 p=ROOT/path; p.parent.mkdir(parents=True,exist_ok=True); p.write_text(text)
def js(path,value): write(path,json.dumps(value,indent=2,ensure_ascii=False)+'\n')
js('data-model/initial-units.json',[dict(code=f'UNIT-{i:02}',name=n,kind='Unclassified',active=True) for i,n in enumerate(UNITS,1)])
# Type notation: text, memo, integer, boolean, date, datetime, money, guid; references below are logical keys, not SharePoint lookups.
entities={
'Events':('Configuration','EventCode:text!,Name:text!,Active:boolean!'),
'EventEditions':('Configuration','EditionCode:text!,EventId:guid!,Year:integer!,Starts:date!,Ends:date!,Status:text!,ClosedAt:datetime'),
'OrganizationalUnits':('Configuration','UnitCode:text!,Name:text!,Kind:text!,Active:boolean!,MergedInto:guid'),
'EditionUnits':('Configuration','EditionId:guid!,UnitId:guid!,DisplayNameSnapshot:text!,OwnerGroupId:integer!,Active:boolean!'),
'Streams':('Configuration','UnitId:guid!,Description:memo'),
'Teams':('Configuration','UnitId:guid!,Responsibility:memo'),
'TeamResponsibilities':('Configuration','EditionId:guid!,TeamId:guid!,UnitId:guid!,Responsibility:text!'),
'Venues':('Operations','EditionId:guid!,Name:text!,Capacity:integer!,Address:memo,Accessible:boolean!'),
'ActivityLocations':('Operations','VenueId:guid!,Name:text!,Capacity:integer!'),
'Sessions':('Operations','EditionId:guid!,LocationId:guid!,StreamId:guid!,Title:text!,Starts:datetime!,Ends:datetime!,CoordinatorId:text!,Capacity:integer!'),
'SessionAssignments':('Operations','SessionId:guid!,ParticipationId:guid!,AttendanceStatus:text!,AttendedAt:datetime'),
'Participants':('RestrictedProfile','FullName:text!,PreferredName:text,Phone:text!,AlternatePhone:text,Email:text,State:text!,District:text,Village:text,Community:memo,Languages:memo,Gender:text,BirthDate:date,IdentityType:text,IdentityStatus:text!'),
'EventParticipations':('Operations','ParticipantId:guid!,EditionId:guid!,PrimaryStreamId:guid!,Category:text!,NominatingTeamId:guid!,OwnerId:text!,Starts:date!,Ends:date!,ReviewStatus:text!,TravelEligible:boolean!,AccommodationEligible:boolean!,FoodEligible:boolean!,KitEligible:boolean!,FinanceEligible:boolean!,ArrivalStatus:text!,RegistrationStatus:text!,DepartureStatus:text!,OpaqueBadgeToken:guid!'),
'ParticipationSnapshots':('Operations','ParticipationId:guid!,NominationRecordKey:guid!,EditionId:guid!,DisplayName:text!,StreamNameSnapshot:text!,CategorySnapshot:text!,ReviewStatus:text!,SourceVersion:integer!'),
'ParticipantStreamAssignments':('Operations','ParticipationId:guid!,StreamId:guid!,IsPrimary:boolean!'),
'Nominations':('StreamPartition','EditionId:guid!,StreamId:guid!,FullName:text!,Phone:text!,State:text!,District:text,Community:memo,Category:text!,ClientRequestKey:guid!,SubmitRequested:boolean!'),
'ReviewQueue':('Reviewer','NominationKey:text!,EditionId:guid!,StreamId:guid!,SourceVersion:text!,Snapshot:memo!,Status:text!,ReviewerId:text,ReviewedAt:datetime,RejectionReason:memo'),
'TravelBookings':('Transport','ParticipationId:guid!,Mode:text!,Origin:text!,ArrivalAt:datetime,DepartureAt:datetime,ServiceNumber:text,PickupPoint:text,BookingReference:text,Status:text!'),
'Vehicles':('Transport','EditionId:guid!,Registration:text!,Capacity:integer!,DriverName:text!,DriverPhone:text!,CoordinatorId:text!'),
'Trips':('Transport','EditionId:guid!,VehicleId:guid!,Origin:text!,Destination:text!,Starts:datetime!,Ends:datetime!,Status:text!'),
'TripPassengers':('Transport','TripId:guid!,ParticipationId:guid!,Status:text!'),
'AccommodationProperties':('Accommodation','EditionId:guid!,Name:text!,Address:memo!,Contact:text'),
'Buildings':('Accommodation','PropertyId:guid!,Name:text!'),
'Rooms':('Accommodation','BuildingId:guid!,Name:text!,Capacity:integer!,Restriction:text!,Accessible:boolean!'),
'Beds':('Accommodation','RoomId:guid!,EditionId:guid!,Label:text!,Active:boolean!'),
'AccommodationAllocations':('Accommodation','ParticipationId:guid!,BedId:guid!,Starts:datetime!,Ends:datetime!,Status:text!,CheckedInAt:datetime,CheckedOutAt:datetime,AllocationKey:text!'),
'MealLocations':('Food','EditionId:guid!,Name:text!,Capacity:integer!'),
'FoodStalls':('Food','LocationId:guid!,Name:text!,Cuisine:text,Capacity:integer!'),
'MealWindows':('Food','EditionId:guid!,LocationId:guid!,Starts:datetime!,Ends:datetime!,Capacity:integer!'),
'MealEntitlements':('Food','ParticipationId:guid!,WindowId:guid!,Quantity:integer!,DietaryInstructions:memo'),
'MealTransactions':('Food','ParticipationId:guid!,WindowId:guid!,Quantity:integer!,ServedAt:datetime!,IssueKey:text!'),
'KitInventory':('Kits','EditionId:guid!,Kind:text!,Movement:text!,Quantity:integer!,Reference:text!,RecordedAt:datetime!'),
'KitAllocations':('Kits','ParticipationId:guid!,Kind:text!,Quantity:integer!'),
'KitIssues':('Kits','ParticipationId:guid!,Kind:text!,IssueKey:text!,IssuedAt:datetime!,Reason:text!,ReplacesIssueId:guid'),
'IdentityVerification':('Identity','ParticipationId:guid!,DocumentType:text!,Status:text!,DocumentReference:text,VerifiedBy:text,VerifiedAt:datetime'),
'EmergencyAssistance':('Safety','ParticipationId:guid!,EmergencyContact:text,AccessibilityInstructions:memo,AssistanceInstructions:memo,ReviewAt:date!'),
'RestrictedBankDetails':('FinanceVault','ParticipantId:guid!,AccountHolder:text!,BankName:text!,AccountNumber:text!,IFSC:text!,VerificationStatus:text!,VerifiedBy:text,VerifiedAt:datetime,ReviewAt:date!'),
'FinancialRequests':('Finance','ParticipationId:guid!,Category:text!,EligiblePaise:integer!,RequestedPaise:integer!,ApprovedPaise:integer!,MakerId:text!,Status:text!'),
'FinancialApprovals':('FinanceChecker','RequestId:guid!,Stage:text!,Decision:text!,ApproverId:text!,DecidedAt:datetime!,RequestVersion:text!'),
'PaymentRecords':('FinancePayer','RequestId:guid!,AmountPaise:integer!,Status:text!,ExternalReference:text,RecordedAt:datetime!,RecordedBy:text!,ReconciledAt:datetime'),
'BadgeIssues':('Registration','ParticipationId:guid!,OpaqueToken:guid!,IssuedAt:datetime!,Reason:text!,RevokedAt:datetime,ReplacesId:guid'),
'UserRoles':('SecurityAdmin','EntraObjectId:text!,RoleId:text!,EditionId:guid!,StreamId:guid,ValidUntil:datetime!'),
'RolePermissions':('SecurityAdmin','RoleId:text!,Resource:text!,CanView:boolean!,CanCreate:boolean!,CanEdit:boolean!,CanApprove:boolean!,CanExport:boolean!,CanDelete:boolean!'),
'HelpdeskTickets':('Helpdesk','EditionId:guid!,ParticipationId:guid,Category:text!,Priority:text!,OwnerId:text!,Status:text!,Summary:memo!,DueAt:datetime,ResolvedAt:datetime'),
'RestrictedIncidents':('Safety','EditionId:guid!,ParticipationId:guid,Category:text!,Priority:text!,OwnerId:text!,Status:text!,Details:memo!,ReviewAt:date!'),
'ApprovalHistory':('Audit','ParticipationId:guid!,Workflow:text!,Stage:text!,ActorObjectId:text!,Decision:text!,At:datetime!,RequestKey:text!,Notes:memo'),
'AuditEvents':('Audit','EditionId:guid!,EntityName:text!,EntityId:guid!,Action:text!,ActorObjectId:text!,At:datetime!,CorrelationId:text!,Outcome:text!'),
'CommandQueue':('Automation','EditionId:guid!,EntityId:guid!,Action:text!,ExpectedVersion:text!,RequestKey:text!,Status:text!,Attempt:integer!,NextAttemptAt:datetime,ErrorCode:text'),
'WorkflowReceipts':('Automation','RequestKey:text!,Stage:text!,State:text!,TargetId:text,SourceVersion:text!,LeaseUntil:datetime,LastErrorCode:text'),
'NotificationOutbox':('Automation','DedupeKey:text!,Template:text!,RecipientObjectId:text!,State:text!,Attempts:integer!,SentAt:datetime')}
retention={'Configuration':'Keep event configuration and immutable yearly labels; annual review.', 'FinanceVault':'Policy decision required before use; delete after approved finance/legal retention and hold review.', 'Identity':'Do not collect document copies by default. Verification status preferred; delete reference after approved purpose ends.', 'Safety':'Short purpose-bound period approved by privacy lead; review immediately after event.', 'RestrictedProfile':'Review after event; retain reusable profile only under documented purpose and notice.', 'Audit':'Approved audit schedule; lock exports and hold exceptions.'}
spec=[]
for name,(boundary,fields) in entities.items():
 fs=[dict(name='RecordKey',type='guid',required=True,unique=True,index=True),dict(name='CreatedAt',type='datetime',required=True),dict(name='UpdatedAt',type='datetime',required=True)]
 for f in fields.split(','):
  key,typ=f.split(':'); required=typ.endswith('!'); typ=typ.rstrip('!')
  fs.append(dict(name=key,type=typ,required=required,index=(key.endswith('Id') or key in ['NominationRecordKey','EditionCode','Status','RequestKey','IssueKey','AllocationKey','NominationKey','DedupeKey','ClientRequestKey']),unique=key in ['EditionCode','RequestKey','IssueKey','AllocationKey','NominationKey','DedupeKey','ClientRequestKey']))
 spec.append(dict(name=name,boundary=boundary,implementation='design-only unless listed in sharepoint/core-lists.json',fields=fs,retention=retention.get(boundary,'Review after event; duration requires approved records schedule and hold policy.'),validation='GUID references must exist; edition references must agree; immutable annual snapshots; domain rules in requirements/invariants.md.'))
js('data-model/entities.json',spec)
md=['# Data dictionary','Logical design: **not all entities are provisioned or implemented**. RecordKey is the application GUID; SharePoint also has its own numeric ID. Cross-list references require service validation. No cascaded deletion.','']
for e in spec:
 md += [f'## {e["name"]}',f'Boundary: {e["boundary"]}. Retention: {e["retention"]}', '', '| Field | Type | Required | Indexed | Unique |','|---|---|---|---|---|']
 for f in e['fields']: md.append('| '+ ' | '.join(str(f.get(k,False)) for k in ['name','type','required','index','unique'])+' |')
 md += ['',e['validation'],'']
write('data-model/dictionary.md','\n'.join(md))
# Explicit conservative default-deny matrix. Grants describe target permissions, not live tenant enforcement.
roles={'Super Administrator':{'Configuration':'vcea','SecurityAdmin':'vcea'},'Event Director / Leadership':{'Dashboard':'v'},'Stream Owner':{'OwnStreamNominations':'vce','OwnStreamProjection':'v'},'Registration Team':{'Registration':'vcea','OperationalProjection':'v','IdentityStatus':'v'},'Accommodation Team':{'Accommodation':'vcea','ArrivalCommands':'vc','OperationalProjection':'v'},'Transport / Mobility Team':{'Transport':'vce','OperationalProjection':'v'},'Food Team':{'Food':'vce','OperationalProjection':'v'},'Gifts & Kits Team':{'Kits':'vce','OperationalProjection':'v'},'Finance & Accounts Team':{'Finance':'vce','FinanceVault':'vce'},'Finance Checker':{'FinanceChecker':'vca','Finance':'v'},'Finance Payer':{'FinancePayer':'vc','Finance':'v'},'Venue / Theme Team':{'Sessions':'vce','AssignedProjection':'v'},'Safety / Public Health Team':{'Safety':'vce','OperationalProjection':'v'},'Helpdesk / Volunteers':{'Helpdesk':'vce','HelpdeskProjection':'v'},'Digital & Analytics Team':{'AnalyticsProjection':'v'},'Participant':{'SelfProjection':'v'}}
with (ROOT/'security/role-permissions.csv').open('w',newline='') as f:
 w=csv.writer(f); w.writerow(['role','resource','view','create','edit','approve','export','delete','scope'])
 for role,grants in roles.items():
  for resource,mask in grants.items(): w.writerow([role,resource,*[str(ch in mask).lower() for ch in 'vceaxd'],'own verified identity' if role=='Participant' else 'assigned stream and edition' if role=='Stream Owner' else 'authorized edition'])
# Minimal tenant core. All other schema is design, not silently provisioned.
core_names=['Events','EventEditions','OrganizationalUnits','EditionUnits','Nominations','ReviewQueue','ParticipationSnapshots','AuditEvents','WorkflowReceipts']
core=[e for e in spec if e['name'] in core_names]
js('sharepoint/core-lists.json',core)
headers=['name','phone','state','district','community','edition','stream','category','participant_id']
with (ROOT/'test-data/participant-import-template.csv').open('w',newline='') as f: csv.writer(f).writerow(headers)
try:
 from openpyxl import Workbook
 from openpyxl.worksheet.datavalidation import DataValidation
 wb=Workbook(); ws=wb.active; ws.title='Participants'; ws.append(headers); ws.freeze_panes='A2'; ws.auto_filter.ref='A1:I5001'
 for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=25
 for r in range(2,5002): ws.cell(r,2).number_format='@'
 help=wb.create_sheet('Instructions'); help.append(['SYNTHETIC DEVELOPMENT TEMPLATE']); help.append(['Phone must be text; do not enter banking, identity, health or emergency data.']); help.append(['Use configured edition and stream GUIDs. Export Participants as UTF-8 CSV for local validation.']); help.append(['Microsoft import path needs tenant validation; this workbook is not an automated tenant importer.'])
 from openpyxl.worksheet.table import Table,TableStyleInfo
 ws.append(['','','','','','','','','']); t=Table(displayName='ParticipantImport',ref='A1:I2');t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(t)
 wb.save(ROOT/'test-data/participant-import-template.xlsx')
except ImportError: pass
print(f'Generated {len(spec)} entity definitions, {len(UNITS)} exact initial names and import templates.')
